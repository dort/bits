"""Build an editable slide deck, audience PDF, and separate presenter documents.

All audience formats use the same positioned text, shapes, equations, and plots.
The only computation experiment here is the independently checked scalar example.
"""

from pathlib import Path
import io
import json
import math
import re
import sys
import textwrap
from xml.sax.saxutils import escape
import zipfile

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from PIL import Image, ImageDraw, ImageFont
import pymupdf
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, KeepTogether
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches as DocInches, Pt as DocPt, RGBColor as DocRGB

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))
from content import SLIDES, SOURCES, PAPER, REPO, COMMIT

SOURCE_SHORT = {"P": "Paper v1", "E": "ePC paper", "I": "sPC solver", "R": "ePC solver",
                "N": "State types", "W": "Local learning", "G": "ePC guide",
                "D": "Stability report", "T": "ePC tests"}

ASSETS = ROOT / "assets"
ASSETS.mkdir(exist_ok=True)
PAGE_W, PAGE_H = 13.333333, 7.5
NAVY, TEAL, CORAL = "14344B", "087F79", "C55F48"
INK, MUTED, LIGHT = "203A4D", "526878", "F5F5EF"
WHITE, BORDER, PALE, GOLD = "FFFFFF", "DCE3DF", "E5F2EE", "A77322"
FONT_DIR = Path("/usr/share/fonts/truetype/lato")
pdfmetrics.registerFont(TTFont("Lato", str(FONT_DIR / "Lato-Regular.ttf")))
pdfmetrics.registerFont(TTFont("Lato-Bold", str(FONT_DIR / "Lato-Bold.ttf")))
pdfmetrics.registerFont(TTFont("Guide", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
pdfmetrics.registerFont(TTFont("Guide-Bold", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"))
plt.rcParams.update({"font.family": "DejaVu Sans", "mathtext.fontset": "dejavusans",
                     "axes.spines.top": False, "axes.spines.right": False})


def rgb(value):
    return RGBColor.from_string(value)


def cpdf(value):
    return colors.HexColor("#" + value)


def wrap(text, width, size, bold=False):
    font = "Lato-Bold" if bold else "Lato"
    lines = []
    for paragraph in text.split("\n"):
        words = paragraph.split()
        if not words:
            lines.append("")
            continue
        line = words.pop(0)
        for word in words:
            if pdfmetrics.stringWidth(line + " " + word, font, size) <= width * 72:
                line += " " + word
            else:
                lines.append(line)
                line = word
        lines.append(line)
    return lines


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(PAGE_W), Inches(PAGE_H)
        self.prs.core_properties.title = "Augmented Lagrangian Predictive Coding — FabricPC"
        self.prs.core_properties.subject = "30-minute research presentation with private speaker notes"
        self.prs.core_properties.author = "FabricPC research presentation"
        self.pdf = canvas.Canvas(str(ROOT / "ALPC_FabricPC_slides.pdf"),
                                 pagesize=(PAGE_W * 72, PAGE_H * 72))
        self.pdf.setTitle("Augmented Lagrangian Predictive Coding — audience slides")
        self.bounds = []

    def new(self, index):
        self.index = index
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.rect(0, 0, PAGE_W, PAGE_H, LIGHT)
        self.rect(0, 0, PAGE_W, .10, TEAL)
        s = SLIDES[index]
        self.text(s["section"], .5, .27, 11.8, 11, TEAL, True)
        if index != 0:
            self.text(s["title"], .5, .68, 12.25, 31, NAVY, True, max_lines=2)
            self.text(s["subtitle"], .52, 1.68, 12.2, 17, MUTED, max_lines=2)
        self.line(.5, 7.0, 12.83, 7.0, BORDER, .8)
        source_codes = s["sources"]
        if source_codes:
            label = "SOURCES  " + " · ".join(f"[{code}]" for code in source_codes)
            if "P" in source_codes:
                label += "  Seely & Gould, 2026"
            if any(code not in ("P", "E") for code in source_codes):
                label += "  |  FabricPC @ 8406e6a"
            self.text(label, .52, 7.15, 10.7, 10, MUTED,
                      link=SOURCES[source_codes[0]][1])
        else:
            self.text("FABRICPC · RESEARCH DISCUSSION", .52, 7.15, 10.7, 10, MUTED)
        self.text(f"{index+1:02d}" + (" / 20" if index < 20 else " · BACKUP"),
                  11.65, 7.12, 1.1, 12, MUTED)

    def rect(self, x, y, w, h, fill, stroke=None, radius=False):
        shape = self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                             Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
        if stroke:
            shape.line.color.rgb = rgb(stroke); shape.line.width = Pt(.7)
        else:
            shape.line.fill.background()
        if radius:
            shape.adjustments[0] = .09
        self.pdf.setFillColor(cpdf(fill))
        self.pdf.setStrokeColor(cpdf(stroke or fill))
        if radius:
            self.pdf.roundRect(x*72, (PAGE_H-y-h)*72, w*72, h*72, 7, fill=1, stroke=bool(stroke))
        else:
            self.pdf.rect(x*72, (PAGE_H-y-h)*72, w*72, h*72, fill=1, stroke=bool(stroke))
        self.bounds.append((self.index, "shape", x, y, w, h))
        return shape

    def text(self, content, x, y, w, size=22, color=INK, bold=False, max_lines=None, link=None):
        lines = wrap(content, w, size, bold)
        if max_lines is not None and len(lines) > max_lines:
            raise ValueError(f"Slide {self.index+1} text needs {len(lines)} lines, max {max_lines}: {content}")
        leading = size * 1.23
        h = (len(lines)*leading+4)/72
        box = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame; tf.clear(); tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_before = Pt(0); p.space_after = Pt(0); p.line_spacing = Pt(leading)
            run = p.add_run(); run.text = line
            run.font.name = "Lato"; run.font.size = Pt(size); run.font.bold = bold
            run.font.color.rgb = rgb(color)
            if link: run.hyperlink.address = link
            self.pdf.setFont("Lato-Bold" if bold else "Lato", size)
            self.pdf.setFillColor(cpdf(color))
            self.pdf.drawString(x*72, (PAGE_H-y)*72 - size*.86 - i*leading, line)
        if link:
            self.pdf.linkURL(link, (x*72, (PAGE_H-y-h)*72, (x+w)*72, (PAGE_H-y)*72), relative=0)
        self.bounds.append((self.index, "text", x, y, w, h))
        return h

    def line(self, x1, y1, x2, y2, color=TEAL, width=2):
        shape = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                    Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        shape.line.color.rgb = rgb(color); shape.line.width = Pt(width)
        self.pdf.setStrokeColor(cpdf(color)); self.pdf.setLineWidth(width)
        self.pdf.line(x1*72, (PAGE_H-y1)*72, x2*72, (PAGE_H-y2)*72)

    def arrow(self, x1, y1, x2, y2, color=TEAL):
        self.line(x1, y1, x2, y2, color, 2)
        angle = math.atan2(y2-y1, x2-x1)
        for side in (-.45, .45):
            self.line(x2, y2, x2-.13*math.cos(angle+side), y2-.13*math.sin(angle+side), color, 2)

    def picture(self, path, x, y, w, h=None):
        im = Image.open(path)
        ih = w * im.height / im.width if h is None else h
        self.slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(ih))
        self.pdf.drawImage(str(path), x*72, (PAGE_H-y-ih)*72, width=w*72, height=ih*72, mask="auto")
        self.bounds.append((self.index, "image", x, y, w, ih))
        return ih

    def equation(self, formula, x, y, w, color=NAVY, size=26, max_h=.85):
        formula = formula.replace(r"\mathsf T", r"\mathsf{T}").replace(r"\mathcal L", r"\mathcal{L}")
        formula = formula.replace(r"\frac12", r"\frac{1}{2}")
        key = f"eq_{self.index+1}_{len(self.bounds)}.png"
        path = ASSETS/key
        fig = plt.figure(figsize=(1, 1)); fig.patch.set_alpha(0)
        fig.text(0, 0, "$"+formula+"$", fontsize=size, color="#"+color)
        fig.savefig(path, dpi=240, transparent=True, bbox_inches="tight", pad_inches=.04)
        plt.close(fig)
        im = Image.open(path)
        actual_w = min(w, max_h*im.width/im.height, im.width/240)
        self.picture(path, x, y, actual_w)

    def panel(self, x, y, w, h, label, body, accent=TEAL, body_size=21):
        self.rect(x, y, w, h, WHITE, BORDER, True)
        self.rect(x+.2, y+.24, .07, .35, accent)
        self.text(label, x+.42, y+.23, w-.65, 22, accent, True)
        self.text(body, x+.28, y+.98, w-.56, body_size, INK)

    def takeaway(self, content):
        self.rect(.5, 6.18, 12.33, .60, NAVY, radius=True)
        self.text(content, .73, 6.32, 11.9, 18, WHITE, max_lines=1)

    def finish(self):
        s = SLIDES[self.index]
        text = speaker_note(self.index, s)
        self.slide.notes_slide.notes_text_frame.text = text
        self.pdf.showPage()

    def save(self):
        self.prs.save(ROOT/"ALPC_FabricPC_presentation.pptx")
        self.pdf.save()
        for slide_i, kind, x, y, w, h in self.bounds:
            assert x >= -.005 and y >= -.005 and x+w <= PAGE_W+.005 and y+h <= PAGE_H+.005, (
                slide_i+1, kind, x, y, w, h)


def time_range(index):
    start = sum(s["seconds"] for s in SLIDES[:index])
    end = start + SLIDES[index]["seconds"]
    if not SLIDES[index]["seconds"]: return "Optional backup"
    return f"{start//60:02d}:{start%60:02d}–{end//60:02d}:{end%60:02d}"


def speaker_note(index, s):
    parts = [f"SLIDE {index+1} — {s['title'].replace(chr(10),' ')}", time_range(index),
             "TALK TRACK", s["talk"], "FROM FIRST PRINCIPLES / PREPARATION", s["prep"]]
    if s["transition"]: parts += ["TRANSITION", s["transition"]]
    if s["sources"]:
        parts += ["SOURCES", "\n".join(f"[{k}] {SOURCES[k][0]}\n{SOURCES[k][1]}" for k in s["sources"])]
    return "\n\n".join(parts)


def assets():
    doc = pymupdf.open(ROOT.parent/"2605.31022v1.pdf")
    for page, clip, name in [(3,(34,25,576,325),"paper_fig2.png"), (4,(35,25,578,260),"paper_fig3.png")]:
        region = pymupdf.Rect(clip)
        pix = doc[page].get_pixmap(matrix=pymupdf.Matrix(3,3), clip=region, alpha=False)
        pixels = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
        rows, columns = np.where(np.min(pixels[:, :, :3], axis=2) < 235)
        tight = pymupdf.Rect(region.x0+columns.min()/3-4, region.y0+rows.min()/3-4,
                            region.x0+columns.max()/3+4, region.y0+rows.max()/3+4)
        doc[page].get_pixmap(matrix=pymupdf.Matrix(3,3), clip=tight & region, alpha=False).save(ASSETS/name)
    eta, alpha, h0, dual0 = .2, .5, .2, 0.
    pc, alm, dual = [h0], [h0], [dual0]
    for _ in range(45):
        pc.append(pc[-1]-eta*(2*pc[-1]-1.2))
        new = alm[-1]-eta*(2*alm[-1]-1.2+dual[-1])
        alm.append(new); dual.append(dual[-1]+alpha*(new-.2))
    fig, axes = plt.subplots(1,2,figsize=(11.9,2.8),layout="constrained")
    fig.patch.set_facecolor("#"+WHITE)
    axes[0].plot(pc, color="#"+CORAL, lw=2.5, label="sPC state")
    axes[0].plot(alm, color="#"+TEAL, lw=2.5, label="PC-ALM state")
    axes[0].axhline(.2, color="#"+MUTED, ls=":", lw=1)
    axes[0].set(ylabel="Hidden activation h", xlabel="Primal-dual cycles (PC: activity steps)", ylim=(.05,.69))
    axes[0].legend(loc="upper right", frameon=False, fontsize=10)
    axes[1].plot(dual, color="#"+TEAL, lw=2.5, label="Multiplier λ")
    axes[1].axhline(.8, color="#"+NAVY, ls=":", label="−BP activation derivative = 0.8")
    axes[1].set(ylabel="Multiplier λ", xlabel="Primal-dual cycles", ylim=(-.04,1.03))
    axes[1].legend(loc="lower right", frameon=False, fontsize=9)
    for ax in axes:
        ax.grid(axis="y", color="#"+BORDER, linewidth=.7)
        ax.tick_params(labelsize=10)
        ax.set_xlim(0,45)
        ax.set_axisbelow(True)
    fig.savefig(ASSETS/"scalar_dynamics.png", dpi=190, facecolor="white")
    plt.close(fig)
    # Independent closed forms and Algorithm 1 endpoint check.
    x=y=w2=rho=1.; w1=.2
    h_pc=(w2*y+rho*w1*x)/(w2*w2+rho)
    grad_pc=np.array([-rho*(h_pc-w1*x)*x, (w2*h_pc-y)*h_pc])
    grad_bp=np.array([(w2*w1*x-y)*w2*x, (w2*w1*x-y)*w1*x])
    h,lam=w1*x,0.
    for _ in range(199):
        h-=eta*(w2*(w2*h-y)+rho*(h-w1*x)+lam)
        lam+=alpha*(h-w1*x)
    h-=eta*(w2*(w2*h-y)+rho*(h-w1*x)+lam)
    grad_alm=np.array([-(lam+rho*(h-w1*x))*x, (w2*h-y)*h])
    assert np.allclose(grad_pc,[-.4,-.24],atol=1e-12)
    assert np.allclose(grad_bp,[-.8,-.16],atol=1e-12)
    assert np.allclose(grad_alm,grad_bp,atol=1e-12)
    eig=np.linalg.eigvals(np.array([[.6,-.2],[.3,.9]]))
    return dict(pc_hidden=h_pc, alm_hidden=h, alm_multiplier=lam,
                pc_gradient=grad_pc.tolist(), bp_gradient=grad_bp.tolist(),
                alm_gradient=grad_alm.tolist(), spectral_radius=float(np.max(abs(eig))))


def row(d, y, cols, widths, header=False, height=.64):
    x=.52
    for j,(label,w) in enumerate(zip(cols,widths)):
        d.rect(x,y,w-.035,height,NAVY if header else (WHITE if j%2 else PALE), BORDER)
        d.text(label,x+.13,y+.09,w-.27,17 if header else 16,
               WHITE if header else INK, header, max_lines=3)
        x+=w


def build_slides():
    d=Deck()
    for i,s in enumerate(SLIDES):
        d.new(i)
        if i==0:
            d.text(s["title"],.65,1.18,11.9,44,NAVY,True,max_lines=2)
            d.text(s["subtitle"],.68,3.02,11.8,25,MUTED)
            d.text("Jeffrey Seely & Julian Gould · Sakana AI · May 2026",.70,3.66,11.8,19,MUTED)
            for x,label,sub,color in [( .7,"sPC","Relax hidden states",CORAL),(4.85,"ePC","Optimize error coordinates",TEAL),(9.,"PC-ALM","Accumulate local dual credit",NAVY)]:
                d.rect(x,4.65,3.6,1.28,WHITE,BORDER,True)
                d.text(label,x+.2,4.83,3.2,25,color,True)
                d.text(sub,x+.2,5.35,3.2,16,INK)
            d.text("30-MINUTE TEAM PRESENTATION  ·  27 MIN TALK + 3 MIN DISCUSSION",.7,6.4,11.9,12,TEAL,True)
        elif i==1:
            for j,(label,sub) in enumerate([("Input x","Observed data"),("Hidden h₁","Learned features"),("Hidden h₂","Learned features"),("Prediction ŷ","Compared with target y")]):
                x=.6+j*3.18
                d.rect(x,2.38,2.6,1.04,WHITE,BORDER,True)
                d.text(label,x+.2,2.58,2.25,23,TEAL,True)
                d.text(sub,x+.05,3.64,2.8,16,MUTED)
                if j<3: d.arrow(x+2.65,2.90,x+3.08,2.90)
            d.equation(r"h_i=\sigma(W_i h_{i-1}),\qquad \hat y=W_L h_{L-1}",.8,4.32,11.7,size=29)
            d.equation(r"\ell=\frac{1}{2}\|y-\hat y\|^2",.8,5.26,4.6,size=27)
            d.text("Weights W are shared. Hidden states h belong to an example.",5.6,5.13,6.35,21,INK)
            d.takeaway("Predictive coding makes hidden activations adjustable during training.")
        elif i==2:
            d.panel(.55,2.25,4.1,3.63,"The derivative we need","How much would the final loss change if this hidden activation changed?",body_size=23)
            d.text("δᵢ: BP derivative of loss with respect to hᵢ",5.0,2.38,7.2,21,TEAL,True)
            d.equation(r"\delta_{L-1}=W_L^{\mathsf T}(\hat y-y)",5.0,3.13,7.3,size=29)
            d.equation(r"\delta_i=J_{i+1}^{\mathsf T}\delta_{i+1}",5.0,4.04,7.3,size=29)
            d.text("Jᵢ₊₁ is the adjacent layer’s Jacobian: its table of local sensitivities.",5.0,4.95,7.4,21)
            d.takeaway("BP computes these signals in a reverse pass through the composed network.")
        elif i==3:
            d.equation(r"r_i=h_i-\sigma(W_i h_{i-1})",.8,2.27,11.8,size=29)
            d.equation(r"F_{\mathrm{PC}}=\ell+\frac{\rho}{2}\sum_{i=1}^{L-1}\|r_i\|^2",.8,3.10,11.8,size=29,max_h=1.0)
            d.panel(.6,4.45,5.94,1.43,"1  Infer states","Fixed weights; take T steps on h.",body_size=19)
            d.panel(6.75,4.45,5.94,1.43,"2  Learn weights","Fixed final states; use local gradients.",body_size=19)
            d.text("r = state − prediction     ρ > 0 = penalty strength     T = inference steps",6.15,2.52,6.1,19,MUTED)
            d.takeaway("An interior activity gradient depends only on its own and adjacent layer terms.")
        elif i==4:
            d.panel(.55,2.28,5.9,3.6,"A compromise is allowed","Lower the supervised loss by moving a hidden state away from its forward prediction.",accent=CORAL,body_size=23)
            d.text("ONE HIDDEN SCALAR",6.87,2.36,5.4,13,TEAL,True)
            d.equation(r"F_{\mathrm{PC}}(h)=\frac12(1-h)^2+\frac12(h-0.2)^2",6.85,3.0,5.7,size=23)
            d.text("Forward hidden state",6.90,4.00,3.7,21)
            d.text("0.2",11.25,3.9,1.2,31,NAVY,True)
            d.text("Settled PC hidden state",6.90,4.92,3.85,21)
            d.text("0.6",11.25,4.82,1.2,31,CORAL,True)
            d.takeaway("Faster inference alone cannot remove a finite-penalty difference from BP.")
        elif i==5:
            d.equation(r"h_i(\varepsilon)=\sigma(W_i h_{i-1}(\varepsilon))+\varepsilon_i",.8,2.32,11.7,size=30)
            d.equation(r"E(\varepsilon)=F_{\mathrm{PC}}(h(\varepsilon)),\quad \varepsilon\leftarrow\varepsilon-\eta_\varepsilon\nabla_\varepsilon E",.8,3.29,11.8,size=28)
            d.panel(.6,4.48,5.94,1.47,"Forward: errors → states","An early error changes downstream states.",body_size=18)
            d.panel(6.75,4.48,5.94,1.47,"Reverse: energy → error gradients","FabricPC differentiates the full derived graph.",body_size=18)
            d.takeaway("ePC uses global error inference and retains a local final weight update.")
        elif i==6:
            d.panel(.55,2.32,5.9,3.52,"Small error step","From zero hidden errors, the first error gradient is the BP activation derivative.\n\nWeight gradients are BP-like to first order.",body_size=21)
            d.panel(6.7,2.32,6.08,3.52,"Near PC equilibrium","More relaxation solves the PC energy’s tradeoff.\n\nStable convergence does not make that objective equal to BP’s.",accent=CORAL,body_size=21)
            d.takeaway("Compare both ePC regimes; report step size, optimizer, and curvature stability.")
        elif i==7:
            d.equation(r"\mathcal{L}_\rho=\ell+\sum_i\lambda_i^{\mathsf T}r_i+\frac{\rho}{2}\sum_i\|r_i\|^2",.75,2.4,12.0,size=32,max_h=1)
            d.panel(.6,3.92,3.88,1.98,"Task loss ℓ","Fits the observed target.",body_size=21)
            d.panel(4.72,3.92,3.88,1.98,"Multiplier λ","Accumulates each hidden constraint’s violations.",body_size=21)
            d.panel(8.84,3.92,3.88,1.98,"Penalty ρ","Supplies quadratic curvature.",body_size=21)
            d.takeaway("Descend in hidden states h; ascend in sample-specific multipliers λ.")
        elif i==8:
            steps=[("01","Initialize","Forward states; hidden multipliers λ = 0."),
                   ("02","Primal step","All free states descend the augmented objective."),
                   ("03","Dual step","Recompute new-state residuals; λ ← λ + αr."),
                   ("04","Finish and learn","Final primal only; recompute local weight gradients.")]
            for j,(n,label,body) in enumerate(steps):
                y=2.22+j*.89
                d.rect(.58,y,.55,.56,TEAL,radius=True)
                d.text(n,.67,y+.11,.40,17,WHITE,True)
                d.text(label,1.36,y+.06,3.00,22,NAVY,True)
                d.text(body,4.65,y+.08,7.67,20)
            d.text("Repeat steps 02–03 exactly T−1 times. Weights stay fixed until 04.",1.36,5.86,11.0,17,MUTED)
            d.takeaway("The dual step reads NEW residuals. α is a separate rate; each new batch resets λ.")
        elif i==9:
            d.text("COMPOSITE CREDIT",.8,2.25,4.4,14,TEAL,True)
            d.equation(r"c_i=\lambda_i+\rho r_i",.8,2.70,5.0,size=34)
            d.text("λ = accumulated violation\nρr = current mismatch pressure",7.0,2.60,5.5,21,MUTED)
            d.equation(r"\nabla_{h_i}\mathcal L_\rho=c_i-J_{i+1}^{\mathsf T}c_{i+1}",.8,3.91,11.8,size=29)
            d.equation(r"\nabla_{W_i}\mathcal L_\rho=-\left(c_i\odot\sigma'(W_i h_{i-1})\right)h_{i-1}^{\mathsf T}",.8,4.90,11.8,size=28)
            d.takeaway("Use the augmented signal in BOTH inference and learning. Keep the true residual separate.")
        elif i==10:
            for x,label,formula,sub in [(.65,"1  Feasibility",r"r_i=0","States return to forward values."),(4.84,"2  Stationarity",r"\lambda_i=-\delta_i","Multipliers reproduce BP credit."),(9.03,"3  Learning",r"\nabla_\theta\mathcal L_\rho=\nabla_\theta\ell","Local weight gradients equal BP.")]:
                d.rect(x,2.43,3.64,2.73,WHITE,BORDER,True)
                d.text(label,x+.19,2.68,3.3,22,TEAL,True)
                d.equation(formula,x+.20,3.39,3.23,size=27,max_h=.65)
                d.text(sub,x+.2,4.35,3.25,18)
            d.text("Endpoint identity: differentiable feedforward maps.\nConvergence theorem: stable linear dynamics with fixed weights.",.75,5.46,11.9,18,MUTED)
            d.takeaway("An endpoint characterization is not a finite-step or nonlinear convergence guarantee.")
        elif i==11:
            d.text("x = y = 1     w₁ = 0.2     w₂ = 1     ρ = 1     ηₕ = 0.2     α = 0.5",.67,2.17,12,16,MUTED)
            d.picture(ASSETS/"scalar_dynamics.png",.65,2.64,12.0)
            d.text("Weight gradients (w₁, w₂)",.75,5.58,3.6,18,NAVY,True)
            d.text("sPC: (−0.40, −0.24)",4.55,5.58,3.8,19,CORAL,True)
            d.text("ALM = BP: (−0.80, −0.16)",8.50,5.58,4.1,19,TEAL,True)
            d.takeaway("Independent teaching calculation. Traces show repeated cycles; endpoint values are checked.")
        elif i==12:
            d.rect(.55,2.17,9.05,3.92,WHITE,BORDER,True)
            im=Image.open(ASSETS/"paper_fig3.png")
            w=min(8.73,3.78*im.width/im.height)
            d.picture(ASSETS/"paper_fig3.png",.71+(8.73-w)/2,2.23,w)
            d.text("FIGURE 3",9.90,2.35,2.8,14,TEAL,True)
            d.text("PC: credit near the output.\n\nALM: credit reaches earlier layers.\n\nWidth 32; depth 64.\nOne sample at initialization.",9.90,2.91,2.85,18)
            d.takeaway("“Ballistic” describes a wavefront; it does not imply instant credit or complete convergence.")
        elif i==13:
            d.rect(.55,2.18,7.24,3.92,WHITE,BORDER,True)
            im=Image.open(ASSETS/"paper_fig2.png")
            w=min(6.99,3.79*im.width/im.height)
            d.picture(ASSETS/"paper_fig2.png",.66+(6.99-w)/2,2.23,w)
            d.text("FIGURE 2 · THE TESTED GRID",8.16,2.31,4.58,14,TEAL,True)
            d.text("Widths + depths\n8, 16, 32, 64, 128\n\nIdentity, tanh, ReLU\n\nT = 2L activity steps\n\nOne training epoch",8.16,2.84,4.5,20)
            d.takeaway("Reported BP-level performance in this grid. No ePC baseline or long-training result.")
        elif i==14:
            d.panel(.55,2.25,3.91,3.62,"Theory","BP endpoint identity.\n\nConditional linear convergence.\n\nNonlinear convergence remains open.",body_size=20)
            d.panel(4.71,2.25,3.91,3.62,"Evidence","Two small datasets.\n\nResidual MLPs.\n\nOne epoch; no FabricPC or ePC benchmark.",accent=CORAL,body_size=20)
            d.panel(8.87,2.25,3.91,3.62,"Cost","Extra dual tensor and residual work.\n\nCoupled rate tuning.\n\nMeasure actual time and memory.",body_size=20)
            d.takeaway("The paper’s compute gain is versus widening PC for credit alignment, not versus BP or ePC.")
        elif i==15:
            widths=[2.28,3.34,3.34,3.34]
            row(d,2.25,["Property","sPC","ePC in FabricPC","PC-ALM"],widths,True,.58)
            for y,values in zip([2.90,3.62,4.34,5.06],[
                ["Relaxed variables","Hidden states h","Errors ε; derive h","States h + multipliers λ"],
                ["Inference coupling","Adjacent-layer derivatives","Full derived-graph reverse pass","Adjacent primal-dual updates"],
                ["Target endpoint","Soft PC-energy stationary point","Same PC-energy stationary points*","BP at feasible stationary point"],
                ["Main tradeoff","Simple local energy; finite-penalty gap","Fast signal access; global AD + curvature","Local BP credit; extra state + stability"]]):
                row(d,y,values,widths,height=.70)
            d.text("* Compatible acyclic objective and gradients. Cyclic unrolling and nonlinear μPC need separate treatment.",.6,5.88,12.1,12,MUTED)
            d.takeaway("An optimized error ε and an accumulated multiplier λ are different variables with different jobs.")
        elif i==16:
            d.panel(.58,2.26,5.97,1.70,"State contract","Per-example hidden dual state; initialization, reset, pytrees, batching.",body_size=19)
            d.panel(6.79,2.26,5.97,1.70,"Objective + learning","Augmented local term in both activity and weight derivative paths.",body_size=19)
            d.panel(.58,4.22,5.97,1.70,"Inference lifecycle","New-state residuals; exact dual schedule; final primal-only step.",body_size=19)
            d.panel(6.79,4.22,5.97,1.70,"Diagnostics + scope","Residual, dual, credit, stationarity; audit clamps, energies, scaling.",body_size=19)
            d.takeaway("Start with an acyclic Gaussian model and extend the shared contracts consistently.")
        elif i==17:
            d.panel(.55,2.27,3.91,3.64,"1  Algebra","Scalar + linear graphs.\n\nα = 0 reproduces sPC.\n\nSettled stable ALM gradients match BP.",body_size=20)
            d.panel(4.71,2.27,3.91,3.64,"2  Fair baselines","BP, sPC, both ePC regimes, PC-ALM.\n\nMeasure gradient error, time, memory, and stability.",body_size=20)
            d.panel(8.87,2.27,3.91,3.64,"3  Transfer","Reproduce the paper.\n\nThen longer training and CIFAR conv models.\n\nMultiple seeds + tuned rates.",body_size=20)
            d.takeaway("Success should be a measured benefit at comparable cost, with the learning objective stated.")
        elif i==18:
            d.text("Build a bounded research implementation\nand compare it with the ePC we actually use.",.75,2.6,11.8,33,NAVY,True,max_lines=3)
            d.panel(.7,4.25,5.83,1.63,"If locality is the objective","PC-ALM is a compelling mechanism to test.",body_size=20)
            d.panel(6.77,4.25,5.83,1.63,"If GPU speed is the objective","The advantage over ePC is still unmeasured.",body_size=20)
            d.takeaway("A zero residual can coexist with a nonzero learning signal: the multiplier carries the credit.")
        elif i==19:
            for y,n,label in [(2.65,"01","Local dynamics that reproduce BP credit?"),(3.76,"02","Better learning in deep, narrow networks?"),(4.87,"03","Lower time or memory than our current ePC?")]:
                d.text(n,.80,y,.75,25,TEAL,True)
                d.text(label,1.8,y,10.7,28,NAVY,True)
            d.takeaway("Proposed first step: a linear acyclic solver, then one matched deep–narrow experiment.")
        elif i==20:
            pairs=[("hᵢ / Wᵢ","Hidden activation / layer weight matrix"),("rᵢ / εᵢ","Measured residual / optimized ePC error coordinate"),("λᵢ / cᵢ","Multiplier / composite credit λᵢ + ρrᵢ"),("δᵢ / Jᵢ₊₁","BP activation derivative / adjacent layer Jacobian"),("ρ / α","Penalty strength / dual ascent rate"),("ηₕ / ηε / ηθ","Activity / error-coordinate / weight step sizes"),("L / N / T","Depth / width / inference activity-step count")]
            for j,(a,b) in enumerate(pairs):
                d.text(a,.8,2.26+j*.49,2.55,20,TEAL,True)
                d.text(b,3.65,2.26+j*.49,8.7,20)
            d.takeaway("Sign convention: r = state − prediction; at the ALM endpoint λ = −δ.")
        elif i==21:
            d.equation(r"h^*_{\mathrm{PC}}=\frac{w_2 y+\rho w_1x}{w_2^2+\rho}",.8,2.33,11.6,size=30,max_h=1.1)
            d.equation(r"h^*_{\mathrm{ALM}}=w_1x,\qquad\lambda^*=w_2(y-w_2w_1x)",.8,3.66,11.6,size=29)
            d.equation(r"\nabla_{w_1}\mathcal L_\rho=-\lambda^*x,\qquad\nabla_{w_2}\mathcal L_\rho=(w_2w_1x-y)w_1x",.8,4.85,11.6,size=27)
            d.takeaway("At x = y = w₂ = ρ = 1 and w₁ = 0.2: hPC = 0.6; hALM = 0.2; λALM = 0.8.")
        elif i==22:
            d.text("Q = ∂h/∂ε is invertible and block lower triangular on a DAG.",.75,2.3,11.8,23,TEAL,True)
            d.equation(r"\nabla_\varepsilon E=Q^{\mathsf T}\nabla_hF_{\mathrm{PC}}",.8,3.1,11.7,size=30)
            d.equation(r"\Delta h\approx-\eta_\varepsilon QQ^{\mathsf T}\nabla_hF_{\mathrm{PC}}",.8,4.09,11.7,size=30)
            d.text("Same stationary points; changed geometry and trajectories.\nAt a stationary point: Hε = Qᵀ Hh Q.",.8,5.12,11.4,21)
            d.takeaway("Global signal access comes from the error-to-state dependency graph, not a new objective.")
        elif i==23:
            d.text("Linear residual r = Ah + b; readout ŷ = Ch; curvature B = CᵀC.",.75,2.22,11.8,20,TEAL,True)
            d.equation(r"K=I-\eta_h(B+\rho A^{\mathsf T}A)",.8,2.91,11.5,size=27)
            # Four matrix blocks are laid out explicitly to avoid a TeX dependency.
            d.text("M =",.8,3.95,1.2,29,NAVY,True)
            d.rect(2.14,3.59,10.42,1.46,WHITE,BORDER,True)
            for x,y,f in [(2.5,3.77,r"K"),(7.2,3.77,r"-\eta_h A^{\mathsf T}"),(2.5,4.43,r"\alpha A K"),(7.2,4.43,r"I-\alpha\eta_h A A^{\mathsf T}")]:
                d.equation(f,x,y,4.7,size=25,max_h=.48)
            d.text("Require spectral radius spr(M) < 1. If B = 0 only: ηₕs²(2ρ + α) < 4 for every singular value s of A.",.8,5.34,11.8,19)
            d.takeaway("The constraint-only bound is not a universal guarantee for a nonlinear supervised model.")
        elif i==24:
            d.panel(.57,2.25,5.96,3.63,"Preserve mathematical meaning","error remains state − prediction.\n\nDuals belong to examples.\n\nThe supervised loss is not a hidden equality constraint.",body_size=20)
            d.panel(6.78,2.25,5.96,3.63,"Audit all derivative paths","Recomputed residual after primal.\n\nAugmented term in local learning.\n\nScaling, custom energies, cycles, and masks require explicit scope.",body_size=20)
            d.takeaway("A buffer containing composite credit cannot substitute for the augmented objective contract.")
        elif i==25:
            d.text("[P] Seely & Gould · Augmented Lagrangian Predictive Coding",.7,2.2,11.7,22,NAVY,True,link=PAPER)
            d.text("arXiv:2605.31022v1 · supplied paper · Algorithm 1; Appendices A, C–G",.7,2.71,11.7,17,MUTED)
            d.text("[E] Goemaere et al. · ePC: Fast and Deep Predictive Coding in Digital Simulation",.7,3.25,11.7,21,NAVY,True,link=SOURCES["E"][1])
            d.text("arXiv:2505.20137 · implementation comparison grounded in FabricPC",.7,3.76,11.7,17,MUTED)
            d.text("FabricPC source snapshot · 22 September 2026",.7,4.30,11.7,22,TEAL,True,link=REPO)
            d.text(COMMIT,.7,4.83,11.7,17,MUTED)
            d.text("[I] inference   [R] ePC   [N] state types   [W] learning   [G] guide   [D] report   [T] tests",.7,5.36,11.7,16,INK)
            d.takeaway("Full links and derivations are in the private speaker guide. Code and plots are reproducible.")
        d.finish()
    d.save()
    return d


def inline_parts(text):
    """Convert simple mathematical indices to real document sub/superscripts."""
    pattern = re.compile(r"([_^])(?:\{([^}]+)\}|([A-Za-z0-9\u0370-\u03ff]+))")
    offset = 0
    for match in pattern.finditer(text):
        yield None, text[offset:match.start()]
        yield ("sub" if match[1] == "_" else "super"), match[2] or match[3]
        offset = match.end()
    yield None, text[offset:]


def inline_pdf(text):
    return "".join(f"<{kind}>{escape(value)}</{kind}>" if kind else escape(value)
                   for kind, value in inline_parts(text)).replace("\n", "<br/>")


def build_guides():
    md=["# Augmented Lagrangian Predictive Coding — speaker guide",
        "27 minutes of prepared material + 3 minutes of discussion. Slides 21–26 are optional backup.",
        "Use the talk track during the presentation. Preparation paragraphs provide extensive derivations and caveats; do not read them all aloud.",
        "## Running order"]
    for i,s in enumerate(SLIDES): md.append(f"- {i+1:02d}. {time_range(i)} — {s['title'].replace(chr(10),' ')}")
    for i,s in enumerate(SLIDES):
        md += [f"\n## Slide {i+1}: {s['title'].replace(chr(10),' ')}",f"**Timing:** {time_range(i)}",
               "### Talk track",s["talk"],"### From first principles / preparation",s["prep"]]
        if s["transition"]: md += ["### Transition",s["transition"]]
        md += ["### Sources"]+[f"- [{k}] [{SOURCES[k][0]}]({SOURCES[k][1]})" for k in s["sources"]]
    (ROOT/"ALPC_FabricPC_speaker_guide.md").write_text("\n\n".join(md),encoding="utf-8")

    document=Document()
    sec=document.sections[0]
    sec.top_margin=sec.bottom_margin=DocInches(.7)
    sec.left_margin=sec.right_margin=DocInches(.75)
    normal=document.styles["Normal"]
    normal.font.name="DejaVu Sans"; normal.font.size=DocPt(10)
    normal.paragraph_format.space_after=DocPt(7)
    normal.paragraph_format.line_spacing=1.12
    for name in ("Title","Heading 1","Heading 2"):
        document.styles[name].font.name="Lato"
        document.styles[name].font.color.rgb=DocRGB.from_string(NAVY)
    document.add_heading("Augmented Lagrangian\nPredictive Coding",0)
    document.add_paragraph("PRIVATE SPEAKER GUIDE",style="Subtitle")
    document.add_paragraph("27-minute talk + 3-minute discussion • 20 main slides + 6 backup slides")
    document.add_paragraph("The Talk track is the suggested spoken explanation. From first principles / preparation contains additional derivations and caveats for rehearsal and questions. The complete notes are also embedded privately in the PowerPoint. The audience PDF has no speaker notes.")
    document.add_heading("Running order",1)
    for i,s in enumerate(SLIDES):
        document.add_paragraph(f"{i+1:02d}   {time_range(i)}   {s['title'].replace(chr(10),' ')}")
    for i,s in enumerate(SLIDES):
        document.add_page_break()
        document.add_heading(f"{i+1:02d}  {s['title'].replace(chr(10),' ')}",1)
        document.add_paragraph(time_range(i),style="Subtitle")
        for heading,key in [("Talk track","talk"),("From first principles / preparation","prep"),("Transition","transition")]:
            if s[key]:
                document.add_heading(heading,2)
                for para in s[key].split("\n\n"):
                    paragraph = document.add_paragraph()
                    for kind, value in inline_parts(para):
                        run = paragraph.add_run(value)
                        if kind == "sub": run.font.subscript = True
                        if kind == "super": run.font.superscript = True
        if s["sources"]:
            document.add_heading("Sources",2)
            p=document.add_paragraph()
            for k in s["sources"]:
                link=OxmlElement("w:hyperlink")
                relation=p.part.relate_to(SOURCES[k][1], RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
                link.set(qn("r:id"),relation)
                run=OxmlElement("w:r"); props=OxmlElement("w:rPr")
                color=OxmlElement("w:color"); color.set(qn("w:val"),TEAL)
                size=OxmlElement("w:sz"); size.set(qn("w:val"),"17")
                props.append(color); props.append(size); run.append(props)
                label=OxmlElement("w:t"); label.text=f"[{k}] {SOURCE_SHORT[k]}"
                run.append(label); link.append(run); p._p.append(link)
                p.add_run("   ")
    document.core_properties.title="ALPC — private speaker guide"
    document.save(ROOT/"ALPC_FabricPC_speaker_guide.docx")

    styles={
        "title":ParagraphStyle("title",fontName="Guide-Bold",fontSize=27,leading=33,textColor=cpdf(NAVY),spaceAfter=16),
        "h1":ParagraphStyle("h1",fontName="Guide-Bold",fontSize=18,leading=23,textColor=cpdf(NAVY),spaceAfter=10),
        "h2":ParagraphStyle("h2",fontName="Guide-Bold",fontSize=11,leading=15,textColor=cpdf(TEAL),spaceBefore=10,spaceAfter=6,keepWithNext=True),
        "body":ParagraphStyle("body",fontName="Guide",fontSize=9.8,leading=13.4,spaceAfter=7),
        "source":ParagraphStyle("source",fontName="Guide",fontSize=7.9,leading=10.2,spaceAfter=6,textColor=cpdf(MUTED)),
        "toc":ParagraphStyle("toc",fontName="Guide",fontSize=9,leading=12,spaceAfter=6),
    }
    story=[]
    def add(text,style="body"):
        story.append(Paragraph(inline_pdf(text),styles[style]))
    add("Augmented Lagrangian\nPredictive Coding","title")
    add("PRIVATE SPEAKER GUIDE","h2")
    add("27-minute talk + 3-minute discussion\n20 main slides + 6 optional backup slides")
    add("Use Talk track during the allotted time. From first principles / preparation is additional rehearsal and reference material. This guide contains approximately 8,400 words; it is not a script to read in full during a half-hour slot.")
    add("The notes are also embedded in the PowerPoint notes pane. The separate audience PDF contains slides only. Sources are linked after each slide’s notes, with a pinned FabricPC commit for reproducibility.")
    add("Source distinction","h2")
    add("Paper findings are attributed to the supplied Seely–Gould preprint. The scalar example is independently computed. Repository statements come from code and documentation inspection, not new library benchmarks. The FabricPC integration and evaluation sequence are proposals.")
    story.append(PageBreak())
    add("Running order","h1")
    for i,s in enumerate(SLIDES): add(f"{i+1:02d}   {time_range(i)}   {s['title'].replace(chr(10),' ')}","toc")
    for i,s in enumerate(SLIDES):
        story.append(PageBreak())
        add(f"{i+1:02d}  {s['title'].replace(chr(10),' ')}","h1")
        add(time_range(i),"source")
        for heading,key in [("Talk track","talk"),("From first principles / preparation","prep"),("Transition","transition")]:
            if s[key]:
                add(heading,"h2")
                for para in s[key].split("\n\n"): add(para)
        if s["sources"]:
            add("Sources","h2")
            links=[]
            for k in s["sources"]:
                title,url=SOURCES[k]
                links.append(f'[{k}] <a href="{escape(url)}" color="#{TEAL}">{escape(SOURCE_SHORT[k])}</a>')
            story.append(Paragraph(" &nbsp; · &nbsp; ".join(links),styles["source"]))
    def footer(c,doc):
        c.saveState(); c.setFont("Guide",8); c.setFillColor(cpdf(MUTED))
        c.drawString(48,29,"ALPC · PRIVATE SPEAKER GUIDE")
        c.drawRightString(564,29,str(doc.page)); c.restoreState()
    SimpleDocTemplate(str(ROOT/"ALPC_FabricPC_speaker_guide.pdf"),pagesize=(612,792),
                      rightMargin=48,leftMargin=48,topMargin=40,bottomMargin=42,
                      title="ALPC — private speaker guide").build(story,onFirstPage=footer,onLaterPages=footer)


def validate_and_render(numerics):
    slide_pdf=pymupdf.open(ROOT/"ALPC_FabricPC_slides.pdf")
    notes_pdf=pymupdf.open(ROOT/"ALPC_FabricPC_speaker_guide.pdf")
    assert len(slide_pdf)==26
    source_text="\n".join(p.get_text() for p in slide_pdf)
    assert "FROM FIRST PRINCIPLES / PREPARATION" not in source_text
    assert "TALK TRACK" not in source_text
    with zipfile.ZipFile(ROOT/"ALPC_FabricPC_presentation.pptx") as z:
        private_parts=[n for n in z.namelist() if n.startswith("ppt/notesSlides/notesSlide") and n.endswith(".xml")]
        assert len(private_parts)==26
        assert all("FROM FIRST PRINCIPLES / PREPARATION" in z.read(n).decode() for n in private_parts)
        audience_parts=[n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
        assert all("FROM FIRST PRINCIPLES / PREPARATION" not in z.read(n).decode() for n in audience_parts)
    thumbs=[]
    thumb_w=480
    for i,page in enumerate(slide_pdf):
        pix=page.get_pixmap(matrix=pymupdf.Matrix(thumb_w/page.rect.width,thumb_w/page.rect.width))
        thumb=Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")
        thumbs.append(thumb)
        if i in (0,7,8,9,11,12,13,15,16,23):
            page.get_pixmap(matrix=pymupdf.Matrix(1.5,1.5)).save(ASSETS/f"preview_slide_{i+1:02d}.png")
    tile_h=300
    sheet=Image.new("RGB",(4*500,math.ceil(len(thumbs)/4)*tile_h),"#e7e9e5")
    draw=ImageDraw.Draw(sheet)
    font=ImageFont.truetype(str(FONT_DIR/"Lato-Regular.ttf"),14)
    for i,im in enumerate(thumbs):
        x=10+(i%4)*500; y=8+(i//4)*tile_h
        sheet.paste(im,(x,y)); draw.text((x,y+273),f"{i+1:02d}  {SLIDES[i]['section']}",font=font,fill="#"+NAVY)
    sheet.save(ROOT/"slide_contact_sheet.png")
    notes_pdf[3].get_pixmap(matrix=pymupdf.Matrix(1.3,1.3)).save(ASSETS/"preview_speaker_guide.png")
    report={"slides":len(slide_pdf),"main_slides":20,"backup_slides":6,"talk_seconds":1620,
            "discussion_seconds":180,"private_notes_parts":26,"speaker_guide_pages":len(notes_pdf),
            "speaker_guide_words":sum(len((s['talk']+' '+s['prep']).split()) for s in SLIDES),
            "fabricpc_commit":COMMIT,"scalar_verification":numerics,
            "audience_slides_contain_private_note_sections":False,
            "rendering":"PDF and PowerPoint generated from shared positioned primitives; PDF previews inspected separately.",
            "limitations":"No native PowerPoint/LibreOffice renderer used. No new FabricPC or paper training benchmark run."}
    (ROOT/"validation.json").write_text(json.dumps(report,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(report,indent=2))


if __name__=="__main__":
    numerical_checks=assets()
    build_slides()
    build_guides()
    validate_and_render(numerical_checks)
